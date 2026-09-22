#!/usr/bin/env python3
"""Generate a presenter speech script (Word .docx) matching the 25-slide deck."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Taleed_Procurement_Speech_Script.docx")

NAVY = RGBColor(0x0A, 0x1D, 0x5C)
BLUE = RGBColor(0x17, 0x41, 0xC9)
AMBER = RGBColor(0xB0, 0x6A, 0x0E)
SLATE = RGBColor(0x64, 0x74, 0x8B)

# (slide number, slide title, on-screen hint, spoken speech)
SLIDES = [
    (1, "Cover — Prototype Walkthrough Across Three Roles",
     "Navy title slide.",
     "Good morning everyone, and thank you for joining. Today I'm going to walk you through the "
     "Taleed Procurement Self-Assessment prototype. It's a working, local-first web application — "
     "not a mock-up — and I'll show it from the perspective of all three of its users: the company "
     "Champion who fills in the assessment, the Taleed Analyst who reviews results, and the Super "
     "Admin who manages the whole program. Everything you'll see is running live with synthetic "
     "demonstration data."),

    (2, "Introduction — What is the tool?",
     "Four capability-area cards and a fact strip.",
     "So, what is this tool actually for? In simple terms, it helps a company understand how mature "
     "its procurement function is. The company answers forty straightforward Yes or No questions "
     "spread across four capability areas — Category Management, Spend Analysis, Strategic Sourcing, "
     "and Supplier Relationships. The moment they finish, they get a maturity score from zero to one "
     "hundred percent, a maturity level, and a short list of practical next steps. There's no "
     "consultant required and no spreadsheet to wrestle with — just a clear, structured picture of "
     "where they stand and where to focus."),

    (3, "Agenda — What this deck covers",
     "Five-point numbered agenda.",
     "Here's how I'll structure the demo. First, a quick word on the product itself. Then I'll take "
     "you through the three journeys in turn: the Champion completing an assessment end to end, the "
     "Analyst reviewing the portfolio, and the Super Admin managing exports, corrections and access. "
     "Finally, I'll explain exactly how the scoring works so the numbers you see make sense. Let's "
     "start with the company's point of view."),

    (4, "Section — The Company Champion",
     "Navy section divider.",
     "This first part is the heart of the product: the Company Champion. This is the authorized "
     "representative at a company — think a Head of Procurement — who actually completes the "
     "assessment. I'll go through their entire journey, from creating an account all the way to "
     "receiving their results and report. And importantly, everything they enter is saved locally in "
     "their own browser as they go."),

    (5, "Champion 2.1 — Public landing",
     "Public landing page.",
     "This is the public landing page — the first thing a visitor sees. It sets expectations clearly: "
     "forty Yes/No questions, four capability areas, and immediate results. The messaging is "
     "deliberately simple and reassuring, because we want procurement leaders to feel this is a quick, "
     "focused diagnostic rather than a heavy audit. From here they start by creating an account."),

    (6, "Champion 2.2 — Register a demo account",
     "Registration form, filled in.",
     "Registration is straightforward — full name, work email, job title, and a password. I want to "
     "flag one thing here: this is a prototype, so passwords are never actually stored, and no real "
     "email is sent. It's a faithful simulation of the sign-up experience. Notice the form validates "
     "as you type, so mistakes are caught early rather than after submitting."),

    (7, "Champion 2.3 — Simulated verification",
     "Email verification screen.",
     "Next is email verification. In a live system this would be a code sent to your inbox; in the "
     "prototype we use a fixed demonstration code. This step is here to show the complete, realistic "
     "flow — a company confirms their email before proceeding. Once verified, they set up their "
     "company profile."),

    (8, "Champion 2.4 — Company profile",
     "Company profile form.",
     "Here the Champion tells us about their organization — company name, country, size, and an "
     "authority declaration confirming they're allowed to respond on the company's behalf. This "
     "profile is what links the assessment to the right organization later on, when the Taleed team "
     "looks across the whole portfolio. We also prevent duplicate companies being created by accident."),

    (9, "Champion 2.5 — Champion dashboard",
     "Company dashboard.",
     "This is the Champion's home base — their dashboard. It shows a clear next step, tracks their "
     "completion, and breaks the assessment into the four capability-area cards. One deliberate design "
     "choice: the percentage here is completion — how much they've filled in — not their maturity "
     "score. We keep those two ideas separate so nobody mistakes progress for performance. Let's go "
     "in and answer some questions."),

    (10, "Champion 2.6 — Answering questions",
     "Assessment question view.",
     "This is the core interaction. Each question is a simple Yes, No, or left unanswered — and 'No' "
     "genuinely means the practice isn't consistently in place, it's not a penalty. Every answer saves "
     "instantly to the browser, so there's no 'save' button to remember and nothing is lost if they "
     "step away. They can move between the four sections freely and come back later."),

    (11, "Champion 2.7 — Review before submit",
     "Review screen.",
     "Before submitting, the Champion gets a full review of all forty answers. If anything's missing, "
     "the tool highlights it and links them straight to the unanswered question. This guarantees a "
     "complete, honest picture — you can only submit once all forty are answered. It's a small step "
     "that protects the quality of the result."),

    (12, "Champion 2.8 — Instant results",
     "Results overview with score ring.",
     "And here's the payoff — instant results. The moment they submit, they see their overall maturity "
     "score, a breakdown across the four domains, their maturity band, and their three relative focus "
     "areas. There's no waiting, no manual calculation, no approval delay. This single screen answers "
     "the question every procurement leader has: 'Where do we actually stand?'"),

    (13, "Champion 2.9 — 16 recommendations",
     "Recommendations grid.",
     "Results are only useful if they lead to action, so the tool provides sixteen tailored "
     "recommendations — four for each domain. Crucially, these are selected based on each domain's own "
     "maturity level, so a company that's strong in sourcing but weak in spend analysis gets advice "
     "matched to each. These come straight from the source framework — they're not generic tips or "
     "AI-generated filler."),

    (14, "Champion 2.10 — Printable report",
     "A4 report preview.",
     "For sharing with a wider team or leadership, there's a clean, printable report. It's formatted "
     "for A4 and works with the browser's Print or Save-as-PDF. So the Champion can walk out of this "
     "with a professional summary document to take into a management conversation."),

    (15, "Champion 2.11 — Assessment history",
     "History table.",
     "Finally for the Champion — history. Every submitted assessment is kept. If a correction is ever "
     "needed later, it creates a new revision rather than overwriting the original, so there's always "
     "a full, trustworthy record of how the company's maturity has changed over time. That completes "
     "the company journey — now let's switch to the Taleed side."),

    (16, "Section — The Taleed Analyst",
     "Navy section divider.",
     "The second role is the Taleed Analyst. This is a member of the program team who looks across all "
     "the participating companies. The key thing to understand about the Analyst is that it's a "
     "read-only, oversight role — they can see results and trends, but they can't see companies' "
     "in-progress draft answers, and by default they can't export data."),

    (17, "Analyst 3.1 — Portfolio overview",
     "Portfolio dashboard.",
     "This is the portfolio overview — the Analyst's main screen. It shows participation across the "
     "program, the average maturity, and how companies are distributed across the maturity bands. "
     "Notice the export buttons here are greyed out — that's intentional. An Analyst has visibility, "
     "but exporting data is a permission that has to be granted separately, which we'll see the Super "
     "Admin control shortly."),

    (18, "Analyst 3.2 — Organizations directory",
     "Organizations list.",
     "The organizations directory gives the Analyst a clean list of every participating company, with "
     "their profile details and participation status at a glance. It's the starting point for drilling "
     "into any single company's submitted result."),

    (19, "Analyst 3.3 — Like-for-like comparison",
     "Comparison view.",
     "One of the more powerful features is like-for-like comparison. The Analyst can select two to "
     "four companies and see their results side by side. This is how the program team spots patterns — "
     "which capability areas are consistently strong or weak across the cohort — always on a fair, "
     "comparable basis. Now let's look at the most privileged role: the Super Admin."),

    (20, "Section — The Super Admin",
     "Navy section divider.",
     "The final role is the Super Admin — the Taleed program administrator. They have everything the "
     "Analyst can do, plus the controls that actually run the program: exporting data, managing "
     "corrections, and controlling who has access to what."),

    (21, "Admin 4.1 — Portfolio with exports",
     "Portfolio with export buttons active.",
     "This is the same portfolio the Analyst saw — but look at the difference: for the Super Admin, "
     "the CSV and Excel export buttons are now active. They can pull the filtered results into a real "
     "spreadsheet for reporting or deeper analysis. It's the same data, with an elevated permission."),

    (22, "Admin 4.2 — Organization detail & correction",
     "Organization detail with correction action.",
     "From an organization's detail page, the Super Admin can open a controlled correction. This is "
     "important: it lets a company fix a genuine error in a past submission — but the original "
     "submission stays in place and remains the effective result until the corrected version is "
     "submitted. Nothing is ever silently overwritten, which keeps the whole dataset trustworthy."),

    (23, "Admin 4.3 — People & access",
     "People & access management.",
     "And this is where access is managed. The Super Admin can enable or pause staff accounts and — "
     "remember those greyed-out export buttons? — this is where they'd grant an Analyst permission to "
     "export. There are also sensible guardrails: for example, you can't lock out the last remaining "
     "administrator. That wraps up the three roles."),

    (24, "How scoring works",
     "Four maturity bands.",
     "Let me quickly demystify the scoring, because it's deliberately simple and transparent. Forty "
     "equally weighted questions: every Yes is worth one point, every No is zero. The overall score is "
     "just total Yes answers divided by forty. That maps onto four bands — Foundational, Developing, "
     "Advanced, and Best-in-Class. There's no hidden weighting and no black box. And to be clear, "
     "these are self-reported results for a capability conversation — not a certification or an "
     "external market benchmark."),

    (25, "Closing — Thank you",
     "Navy closing slide.",
     "So to sum up: in a few minutes, a company can go from signing up to a clear, scored, actionable "
     "view of its procurement maturity — and the Taleed team gets a consistent, comparable picture "
     "across every participant. It's fast, it's transparent, and it turns a complex assessment into a "
     "simple conversation. Thank you — I'm very happy to take any questions, or to run through any part "
     "of the live tool in more detail."),
]


def main():
    doc = Document()

    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Title block
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("Taleed Procurement Self-Assessment")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY

    s = doc.add_paragraph()
    r = s.add_run("Presenter Speech Script  ·  25-slide walkthrough")
    r.font.size = Pt(12)
    r.font.color.rgb = SLATE

    hint = doc.add_paragraph()
    r = hint.add_run("Tip: each entry matches one slide, in order. The short italic line is what is on "
                     "screen; the paragraph below it is your spoken script. Aim for a relaxed pace — the "
                     "full deck runs about 8–10 minutes.")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = SLATE

    doc.add_paragraph()

    for no, title, on_screen, speech in SLIDES:
        h = doc.add_paragraph()
        rr = h.add_run(f"Slide {no}  —  {title}")
        rr.bold = True
        rr.font.size = Pt(13)
        rr.font.color.rgb = BLUE
        h.space_before = Pt(10)

        oscreen = doc.add_paragraph()
        rr = oscreen.add_run(f"On screen: {on_screen}")
        rr.italic = True
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = AMBER

        p = doc.add_paragraph(speech)
        p.paragraph_format.space_after = Pt(6)

        # subtle divider
        d = doc.add_paragraph()
        dr = d.add_run("— " * 22)
        dr.font.size = Pt(8)
        dr.font.color.rgb = RGBColor(0xC7, 0xD0, 0xDE)

    doc.save(OUT)
    print("Saved:", OUT, "|", len(SLIDES), "slide scripts")


if __name__ == "__main__":
    main()
