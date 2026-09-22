#!/usr/bin/env python3
"""Build a branded PowerPoint walkthrough of the Taleed Procurement prototype."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "screenshots")
OUT = os.path.join(HERE, "Taleed_Procurement_Walkthrough.pptx")

# Brand palette
NAVY = RGBColor(0x0A, 0x1D, 0x5C)
BLUE = RGBColor(0x17, 0x41, 0xC9)
AMBER = RGBColor(0xF0, 0xA9, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x64, 0x74, 0x8B)
INK = RGBColor(0x1E, 0x29, 0x3B)
BG = RGBColor(0xEE, 0xF1, 0xF5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(1, l, t, w, h)  # rectangle
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(slide, l, t, w, h, s, size, color, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font="Arial"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for i, line in enumerate(s.split("\n")):
        run = p.add_run() if i == 0 else tf.add_paragraph().add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return tb


def title_slide(kicker, title, subtitle):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY)
    box(s, 0, 0, Inches(0.35), SH, BLUE)
    box(s, Inches(0.9), Inches(2.35), Inches(1.5), Inches(0.09), AMBER)
    text(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.5), kicker, 15, AMBER, bold=True)
    text(s, Inches(0.85), Inches(2.6), Inches(11.6), Inches(2.2), title, 46, WHITE, bold=True)
    text(s, Inches(0.9), Inches(4.9), Inches(11), Inches(1.2), subtitle, 18, RGBColor(0xC7, 0xD2, 0xF0))
    text(s, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.5),
         "Local-first React + Redux prototype  ·  Synthetic demonstration data only", 12,
         RGBColor(0x9A, 0xAA, 0xD0))
    return s


def section_slide(no, title, desc):
    s = prs.slides.add_slide(BLANK)
    bg(s, BG)
    box(s, 0, 0, SW, Inches(1.4), NAVY)
    box(s, 0, Inches(1.4), SW, Inches(0.08), AMBER)
    text(s, Inches(0.7), Inches(0.32), Inches(2), Inches(0.8), no, 40, AMBER, bold=True)
    text(s, Inches(2.1), Inches(0.28), Inches(10), Inches(0.9), title, 30, WHITE, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(3), desc, 20, INK)
    return s


def shot_slide(no, title, explanation, filename):
    s = prs.slides.add_slide(BLANK)
    bg(s, BG)
    # Header band
    box(s, 0, 0, SW, Inches(0.92), NAVY)
    box(s, 0, Inches(0.92), SW, Inches(0.06), AMBER)
    text(s, Inches(0.55), Inches(0.06), Inches(1.6), Inches(0.8), no, 22, AMBER, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.75), Inches(0.04), Inches(11), Inches(0.8), title, 20, WHITE, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    # Explanation panel: white card with an amber accent bar and a small label
    exp_t = Inches(1.12)
    exp_h = Inches(0.92)
    box(s, Inches(0.55), exp_t, Inches(12.23), exp_h, WHITE)
    box(s, Inches(0.55), exp_t, Inches(0.09), exp_h, AMBER)
    text(s, Inches(0.85), exp_t + Inches(0.06), Inches(2.2), Inches(0.3), "WHAT YOU SEE", 9, AMBER, bold=True)
    text(s, Inches(0.85), exp_t + Inches(0.30), Inches(11.6), exp_h - Inches(0.36),
         explanation, 13, INK)
    # Image (fit within remaining frame, keep aspect ratio)
    path = os.path.join(SHOTS, filename)
    iw, ih = Image.open(path).size
    frame_l, frame_t = Inches(0.55), Inches(2.25)
    frame_w, frame_h = Inches(12.23), Inches(4.95)
    scale = min(frame_w / iw, frame_h / ih)
    w = int(iw * scale)
    h = int(ih * scale)
    l = int(frame_l + (frame_w - w) / 2)
    t = int(frame_t + (frame_h - h) / 2)
    # subtle white card behind the image
    box(s, l - Emu(38100), t - Emu(38100), w + Emu(76200), h + Emu(76200), WHITE)
    s.shapes.add_picture(path, l, t, width=w, height=h)
    return s


def intro_slide():
    """A single 'what is this tool' overview slide."""
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    box(s, 0, 0, Inches(0.35), SH, BLUE)
    text(s, Inches(0.9), Inches(0.5), Inches(11.6), Inches(0.7),
         "What is the Taleed Procurement Self-Assessment?", 28, NAVY, bold=True)
    box(s, Inches(0.95), Inches(1.32), Inches(1.4), Inches(0.08), AMBER)
    text(s, Inches(0.9), Inches(1.6), Inches(11.6), Inches(1.3),
         "A local-first web tool that helps a company measure how mature its procurement function is. "
         "The company answers 40 simple Yes/No questions across four capability areas and instantly receives "
         "a maturity score, a maturity level and a short list of practical next steps.", 15, INK)
    areas = [
        ("01", "Category\nManagement", "Category ownership, strategies and performance reviews."),
        ("02", "Spend\nAnalysis", "Spend visibility, data quality and decision-ready insight."),
        ("03", "Strategic\nSourcing", "Sourcing discipline, supplier selection and realized value."),
        ("04", "Supplier\nRelationships", "Governance, performance and supplier collaboration."),
    ]
    cw = Inches(2.85)
    gap = Inches(0.25)
    x = Inches(0.9)
    top = Inches(3.05)
    for no, name, desc in areas:
        box(s, x, top, cw, Inches(2.4), BG)
        box(s, x, top, cw, Inches(0.12), BLUE)
        text(s, x + Inches(0.2), top + Inches(0.25), cw - Inches(0.4), Inches(0.6), no, 20, AMBER, bold=True)
        text(s, x + Inches(0.2), top + Inches(0.72), cw - Inches(0.4), Inches(0.9), name, 15, NAVY, bold=True)
        text(s, x + Inches(0.2), top + Inches(1.6), cw - Inches(0.4), Inches(0.75), desc, 11, SLATE)
        x = Emu(int(x) + int(cw) + int(gap))
    box(s, Inches(0.9), Inches(5.9), Inches(11.53), Inches(0.95), NAVY)
    text(s, Inches(1.2), Inches(5.98), Inches(11), Inches(0.8),
         "40 Yes/No questions  \u00b7  4 capability areas  \u00b7  0\u2013100% maturity score  \u00b7  "
         "16 tailored recommendations  \u00b7  3 roles: Champion, Analyst, Super Admin",
         13, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    return s



# ---- Cover
title_slide("TALEED PROCUREMENT SELF-ASSESSMENT",
            "Prototype Walkthrough\nAcross Three Roles",
            "Champion · Analyst · Super Admin — an end-to-end capability diagnostic\ncaptured from a live run of the local prototype.")

# ---- Introduction: what the tool is
intro_slide()

# ---- Agenda
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
box(s, 0, 0, Inches(0.35), SH, BLUE)
text(s, Inches(0.9), Inches(0.6), Inches(11), Inches(0.8), "What this deck covers", 32, NAVY, bold=True)
box(s, Inches(0.95), Inches(1.5), Inches(1.4), Inches(0.08), AMBER)
agenda = [
    ("01", "The product", "A 40-question, 4-domain procurement maturity diagnostic scored 0-100%."),
    ("02", "Champion journey", "Register - verify - profile - answer 40 - submit - results - report - history."),
    ("03", "Analyst journey", "Read-only portfolio, organizations and like-for-like comparison; no exports."),
    ("04", "Super Admin journey", "Portfolio with exports, corrections and people & access management."),
    ("05", "How scoring works", "Yes = 1, No = 0; four maturity bands; 16 tailored recommendations."),
]
y = 1.95
for no, h, d in agenda:
    text(s, Inches(0.95), Inches(y), Inches(1.1), Inches(0.6), no, 22, AMBER, bold=True)
    text(s, Inches(2.0), Inches(y - 0.05), Inches(3.4), Inches(0.6), h, 18, NAVY, bold=True)
    text(s, Inches(5.5), Inches(y - 0.02), Inches(7.3), Inches(0.7), d, 14, SLATE)
    y += 0.95

# ---- Section + shots: Champion
section_slide("02", "The Company Champion",
              "A company's authorized representative registers, completes the assessment and receives an "
              "instant maturity result with tailored recommendations. Everything is saved locally in the browser.")
champ = [
    ("Public landing", "The introduction explains the diagnostic: 40 Yes/No questions, four capability areas, immediate results.", "01-landing.png"),
    ("Register a demo account", "Name, work email, job title and a dummy password. Passwords are never stored — this is a simulation.", "02-register-filled.png"),
    ("Simulated verification", "A fixed demo code (123456) stands in for an email link. No real email is sent.", "03-verify.png"),
    ("Company profile", "Company name, country, size and an authority declaration connect the assessment to the organization.", "04-company-profile.png"),
    ("Champion dashboard", "A clear next step, completion tracking and four capability-area cards. Completion is not a score.", "05-champion-dashboard.png"),
    ("Answering questions", "Each question is Yes / No / unanswered. Answers save instantly to the browser.", "06-answering-questions.png"),
    ("Review before submit", "All 40 answers are checked for completeness, with links to any missing question.", "07-review-answers.png"),
    ("Instant results", "Overall maturity, four domain scores, a band and three relative focus areas.", "08-champion-results.png"),
    ("16 recommendations", "Four tailored actions per domain, selected from that domain's own maturity band.", "09-champion-recommendations.png"),
    ("Printable report", "An A4 report preview for browser Print / Save as PDF.", "10-champion-report.png"),
    ("Assessment history", "Every submitted revision is retained; corrections never overwrite the original.", "11-champion-history.png"),
]
for i, (t, c, f) in enumerate(champ, 1):
    shot_slide(f"2.{i}", t, c, f)

# ---- Section + shots: Analyst
section_slide("03", "The Taleed Analyst",
              "A program analyst reviews participation and maturity across all submitted companies — "
              "read-only. Draft answers are never exposed, and exports are disabled unless granted.")
analyst = [
    ("Portfolio overview", "Participation, average maturity and score distribution across submitted companies. Export buttons are disabled.", "12-analyst-portfolio.png"),
    ("Organizations directory", "Company profiles and participation at a glance.", "13-analyst-organizations.png"),
    ("Like-for-like comparison", "Compare two to four submitted companies side by side.", "14-analyst-compare.png"),
]
for i, (t, c, f) in enumerate(analyst, 1):
    shot_slide(f"3.{i}", t, c, f)

# ---- Section + shots: Admin
section_slide("04", "The Super Admin",
              "The Taleed program administrator has all analyst capabilities plus data exports, controlled "
              "corrections and people & access management.")
admin = [
    ("Portfolio with exports", "Same portfolio as the analyst, but CSV and Excel exports are enabled.", "15-admin-portfolio.png"),
    ("Organization detail & correction", "Open a controlled correction while the original submission stays effective.", "16-admin-organization-detail.png"),
    ("People & access", "Manage staff access and grant analyst export permission. Guardrails protect the last admin.", "17-admin-access.png"),
]
for i, (t, c, f) in enumerate(admin, 1):
    shot_slide(f"4.{i}", t, c, f)

# ---- Scoring slide
s = prs.slides.add_slide(BLANK)
bg(s, BG)
box(s, 0, 0, SW, Inches(1.4), NAVY)
box(s, 0, Inches(1.4), SW, Inches(0.08), AMBER)
text(s, Inches(0.7), Inches(0.28), Inches(2), Inches(0.9), "05", 40, AMBER, bold=True)
text(s, Inches(2.1), Inches(0.28), Inches(10), Inches(0.9), "How scoring works", 30, WHITE, bold=True,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.9), Inches(1.75), Inches(11.5), Inches(1),
     "40 equally weighted Yes/No questions across 4 domains (10 each).  Yes = 1, No = 0.\n"
     "Overall % = total Yes / 40 x 100.  Domain % = domain Yes / 10 x 100.", 17, INK)
bands = [
    ("Foundational", "0 - 40%", RGBColor(0xC0, 0x3B, 0x2B)),
    ("Developing", "41 - 65%", AMBER),
    ("Advanced", "66 - 80%", BLUE),
    ("Best-in-Class", "81 - 100%", RGBColor(0x1F, 0x8B, 0x4C)),
]
cw = Inches(2.85)
gap = Inches(0.25)
x = Inches(0.9)
for name, rng, color in bands:
    box(s, x, Inches(3.7), cw, Inches(2.0), color)
    text(s, x, Inches(4.05), cw, Inches(0.8), rng, 26, WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, Inches(4.95), cw, Inches(0.7), name, 16, WHITE, bold=True, align=PP_ALIGN.CENTER)
    x = Emu(int(x) + int(cw) + int(gap))
text(s, Inches(0.9), Inches(6.0), Inches(11.5), Inches(1),
     "Results are self-reported and for demonstration only — not a certification or an external market benchmark.",
     13, SLATE)

# ---- Closing
title_slide("THANK YOU",
            "A clearer view of\nprocurement capability.",
            "Live prototype: http://127.0.0.1:5173  ·  Demo accounts use password TaleedDemo!2026\n"
            "Champion: register your own  ·  Analyst: analyst@example.com  ·  Super Admin: admin@example.com")

prs.save(OUT)
print("Saved:", OUT, "|", len(prs.slides._sldIdLst), "slides")
